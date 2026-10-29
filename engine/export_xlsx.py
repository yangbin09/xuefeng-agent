# -*- coding: utf-8 -*-
"""把前端提交的志愿草稿渲染成 .xlsx。

只依赖 openpyxl（项目本来就装了，用于读自定义数据表）。
输出两张表：
  1) 志愿表  —— 按用户排定的投档顺序排列，带档位/位次差/趋势
  2) 数据说明 —— 考生信息、参考年份、分档口径、数据来源、免责声明
"""
import io
import re
import time

from . import config

_BUCKET_CN = {'chong': '冲', 'wen': '稳', 'bao': '保'}
_BUCKET_COLOR = {'冲': 'FFE2E2', '稳': 'E6F0E6', '保': 'E2ECF7'}


def _clean_name(s, fallback='志愿表'):
    s = re.sub(r'[\\/:*?"<>|\r\n\t]', '_', str(s or '')).strip()
    return (s or fallback)[:40]


def suggested_filename(payload):
    base = _clean_name(payload.get('title'), '志愿表')
    return base + '.xlsx'


def build_workbook(payload):
    """payload: {title, province, track, year, rank, score, max, items:[...]}"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    items = payload.get('items') or []
    province = payload.get('province') or ''
    track = payload.get('track') or ''
    year = payload.get('year') or ''
    rank = payload.get('rank') or ''
    score = payload.get('score') or ''
    score_lo = payload.get('scoreLo') or 0
    score_hi = payload.get('scoreHi') or 0
    range_mode = bool(payload.get('rangeMode')) or bool(score_lo and score_hi)
    cap = payload.get('max') or 45

    wb = Workbook()
    ws = wb.active
    ws.title = '志愿表'

    thin = Side(style='thin', color='C8C8C8')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor='2F3B52')
    head_font = Font(bold=True, color='FFFFFF', size=11)

    headers = ['序号', '档位', '院校', '城市', '需考到', '位次', '差几分',
               '近三年', '学费', '可报组', '备注']
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    # 标题行上方留白不要，直接数据；表头冻结
    ws.freeze_panes = 'A2'

    for i, it in enumerate(items, 1):
        bucket = _BUCKET_CN.get(it.get('bucket'), '')
        gap = it.get('gap')
        if gap is None:
            gap = (it.get('rank') or 0) - (rank or 0)
        gap_text = it.get('gapText') or ('%+d（位次差）' % gap if isinstance(gap, int) else '')
        grp = it.get('grpTotal') and ('%s/%s' % (it.get('grpOk') or 0, it.get('grpTotal'))) or ''
        ws.append([
            i, bucket, it.get('school') or '', it.get('city') or '',
            it.get('needScore') or it.get('score'), it.get('rank'),
            gap_text, it.get('zoneText') or '',
            it.get('tuition') if it.get('tuition') is not None else '',
            grp, it.get('note') or '',
        ])
        row = ws.max_row
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).border = border
        ws.cell(row=row, column=1).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=2).alignment = Alignment(horizontal='center')
        color = _BUCKET_COLOR.get(bucket)
        if color:
            ws.cell(row=row, column=2).fill = PatternFill('solid', fgColor=color)
        # 差几分：差=红，富余=绿
        gc = ws.cell(row=row, column=7)
        if isinstance(gap_text, str) and gap_text.startswith('差'):
            gc.font = Font(color='C0392B')
        elif isinstance(gap_text, str) and gap_text.startswith('富余'):
            gc.font = Font(color='2E7D32')

    widths = [6, 6, 26, 10, 9, 10, 13, 12, 10, 10, 26]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---------- 空位补足到 cap 行，便于打印 ----------
    for i in range(len(items) + 1, cap + 1):
        ws.append([i] + [''] * (len(headers) - 1))
        row = ws.max_row
        for c in range(1, len(headers) + 1):
            ws.cell(row=row, column=c).border = border

    # ---------- 说明表 ----------
    ws2 = wb.create_sheet('数据说明')
    cnt = {'冲': 0, '稳': 0, '保': 0}
    for it in items:
        k = _BUCKET_CN.get(it.get('bucket'))
        if k:
            cnt[k] += 1

    rows = [
        ('生成时间', time.strftime('%Y-%m-%d %H:%M:%S')),
        ('考生省份', province),
        ('选科', track),
        ('分数区间', f'{score_lo}–{score_hi} 分' if (score_lo and score_hi) else '未填（单点模式）'),
        ('全省位次', rank),
        ('参考年份', f'{year} 年（{province} {track} 一分一段表）'),
        ('志愿容量', f'{len(items)} / {cap}'),
        ('冲稳保', f"冲 {cnt['冲']} · 稳 {cnt['稳']} · 保 {cnt['保']}"),
        ('', ''),
        ('分档口径',
         '区间模式：校线位次 < 区间上端＝冲；落在区间内＝稳；> 区间下端＝保'
         if range_mode else
         '倍率 = 该校最低录取位次 ÷ 考生位次；冲 0.80~0.95 / 稳 0.95~1.10 / 保 1.10~1.45'),
        ('需考到', '该校（最冷门专业组）的门槛分，同为参考年口径；不等于你想报那个专业的分数'),
        ('差几分', '差 N 分＝需要比区间上限再高 N 分；富余 N 分＝区间下限已高于门槛'),
        ('近三年',
         '把该校 2024~2026 三年的录取位次，按考生同一区间各归一次档再数次数（如「稳2 冲1」）。'
         '只数次数，不给百分比——三年样本撑不出录取概率。'),
        ('学费 / 可报组',
         '专业组数据取自掌上高考招生计划与分专业录取线；可报 N/M 组 = 你的选科能报 M 组里的 N 组'),
        ('数据来源', '一分一段表：各省考试院；校线：掌上高考；专业组：掌上高考招生计划'),
        ('', ''),
        ('免责声明', '数据为公开信息整理，仅供决策参考，不构成志愿填报指导意见。'),
        ('', '最终以各省教育考试院官网公布的当年投档线与招生计划为准。'),
        ('', 'AI 生成内容可能出错，请自行核实。'),
    ]
    ws2.append(['项目', '内容'])
    for c in (1, 2):
        cell = ws2.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal='center')
    for k, v in rows:
        ws2.append([k, v])
    ws2.column_dimensions['A'].width = 14
    ws2.column_dimensions['B'].width = 88
    for r in range(2, ws2.max_row + 1):
        ws2.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical='top')
        ws2.cell(row=r, column=1).font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
